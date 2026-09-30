from typing import Optional
from datetime import date, datetime, timezone, timedelta
MSK = timezone(timedelta(hours=3))

from fastapi import APIRouter, Depends, Request, Form, Query, UploadFile, File
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from sqlalchemy.orm import Session
from jose import jwt, JWTError
import io
from openpyxl import Workbook, load_workbook

from app.database import get_db
from app.models import Terminal, StatusHistory, User
from app.auth import SECRET_KEY, ALGORITHM
from app.schemas import TerminalCreate, TerminalUpdate, StatusChange
from app.templates import templates

router = APIRouter()


def get_user_from_cookie(request: Request, db: Session):
    token = request.cookies.get("access_token")
    if not token:
        return None
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = int(payload.get("sub"))
        return db.query(User).filter(User.id == user_id).first()
    except (JWTError, Exception):
        return None


def format_validation_error(e):
    """Превращает ошибку Pydantic в читаемый текст"""
    try:
        errors = e.errors()
        messages = []
        for err in errors:
            field = '.'.join(str(x) for x in err['loc'])
            msg = err['msg']
            messages.append(f"{field}: {msg}")
        return '; '.join(messages)
    except Exception:
        return str(e)


# ============ СПИСОК ТЕРМИНАЛОВ ============
@router.get("", response_class=HTMLResponse)
async def terminals_list(
    request: Request,
    status: Optional[str] = Query(None),
    serial: Optional[str] = Query(None),
    model: Optional[str] = Query(None),
    box: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(30, ge=10, le=200),
    db: Session = Depends(get_db)
):
    query = db.query(Terminal)

    if status:
        query = query.filter(Terminal.status == status)
    if serial:
        query = query.filter(Terminal.serial_number.ilike(f"%{serial}%"))
    if model:
        query = query.filter(Terminal.model.ilike(f"%{model}%"))
    if box:
        query = query.filter(Terminal.box_number.ilike(f"%{box}%"))

    total = query.count()
    total_pages = (total + per_page - 1) // per_page
    terminals = query.order_by(Terminal.id.desc()).offset((page - 1) * per_page).limit(per_page).all()

    return templates.TemplateResponse(
        "terminals_list.html",
        {
            "request": request,
            "terminals": terminals,
            "total": total,
            "page": page,
            "per_page": per_page,
            "total_pages": total_pages,
            "filters": {
                "status": status,
                "serial": serial,
                "model": model,
                "box": box,
            }
        }
    )


# ============ ЭКСПОРТ ============
@router.get("/export")
async def terminals_export(db: Session = Depends(get_db)):
    terminals = db.query(Terminal).order_by(Terminal.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Терминалы"

    headers = ["ID", "Модель", "Прошивка", "Серийный номер", "Партия",
               "Статус", "Банк", "Тип брака", "Комментарий к браку", "Дата прихода"]
    ws.append(headers)

    for t in terminals:
        ws.append([
            t.id,
            t.model,
            t.firmware_version,
            t.serial_number,
            t.box_number,
            t.status,
            t.bank or '',
            t.defect_type or '',
            t.defect_comment or '',
            t.arrival_date.strftime('%Y-%m-%d') if t.arrival_date else ''
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=terminals.xlsx"}
    )


# ============ ИМПОРТ ============
@router.get("/import", response_class=HTMLResponse)
async def terminal_import_form(request: Request):
    return templates.TemplateResponse("terminal_import.html", {"request": request})


@router.post("/import")
async def terminal_import(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user:
        return RedirectResponse(url="/auth/login", status_code=303)

    if not file.filename.endswith('.xlsx'):
        return templates.TemplateResponse(
            "terminal_import.html",
            {"request": request, "error": "Неверный формат файла. Ожидается .xlsx"}
        )

    try:
        contents = await file.read()
        wb = load_workbook(filename=io.BytesIO(contents), read_only=True)
        ws = wb.active

        imported = 0
        skipped = 0
        errors = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:
                continue

            try:
                model = str(row[0]).strip()
                firmware = str(row[1]).strip() if row[1] else ""
                serial = str(row[2]).strip() if row[2] else ""
                box = str(row[3]).strip() if row[3] else ""

                arrival = row[4]
                if hasattr(arrival, 'strftime'):
                    arrival_str = arrival.strftime('%Y-%m-%d')
                else:
                    arrival_str = str(arrival).strip() if arrival else ""

                if not serial:
                    skipped += 1
                    continue

                # Валидация через Pydantic
                try:
                    data = TerminalCreate(
                        model=model or "Не указано",
                        firmware_version=firmware or "—",
                        serial_number=serial,
                        box_number=box or "—",
                        arrival_date=arrival_str if arrival_str else date.today()
                    )
                except Exception as ve:
                    errors.append(f"Строка {row_idx}: {format_validation_error(ve)}")
                    skipped += 1
                    continue

                # Дубликат
                existing = db.query(Terminal).filter(Terminal.serial_number == data.serial_number).first()
                if existing:
                    skipped += 1
                    continue

                terminal = Terminal(
                    model=data.model,
                    firmware_version=data.firmware_version,
                    serial_number=data.serial_number,
                    box_number=data.box_number,
                    arrival_date=data.arrival_date,
                    status="warehouse"
                )
                db.add(terminal)
                db.flush()

                history = StatusHistory(
                    terminal_id=terminal.id,
                    old_status=None,
                    new_status="warehouse",
                    changed_by=user.id,
                    comment="Импортирован из Excel"
                )
                db.add(history)
                imported += 1

            except Exception as e:
                errors.append(f"Строка {row_idx}: {str(e)}")
                skipped += 1

        db.commit()

        return templates.TemplateResponse(
            "terminal_import.html",
            {
                "request": request,
                "success": f"Импортировано: {imported}, пропущено: {skipped}",
                "errors": errors[:10]
            }
        )

    except Exception as e:
        return templates.TemplateResponse(
            "terminal_import.html",
            {"request": request, "error": f"Ошибка чтения файла: {str(e)}"}
        )


# ============ ДОБАВЛЕНИЕ ============
@router.get("/add", response_class=HTMLResponse)
async def terminal_add_form(request: Request):
    return templates.TemplateResponse("terminal_add.html", {
        "request": request,
        "today": date.today().strftime('%Y-%m-%d')
    })


@router.post("/add")
async def terminal_add(
    request: Request,
    model: str = Form(...),
    firmware_version: str = Form(...),
    serial_number: str = Form(...),
    box_number: str = Form(...),
    arrival_date: str = Form(...),
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user:
        return RedirectResponse(url="/auth/login", status_code=303)

    today = date.today().strftime('%Y-%m-%d')

    # Валидация
    try:
        data = TerminalCreate(
            model=model.strip(),
            firmware_version=firmware_version.strip(),
            serial_number=serial_number.strip(),
            box_number=box_number.strip(),
            arrival_date=arrival_date
        )
    except Exception as e:
        return templates.TemplateResponse("terminal_add.html", {
            "request": request,
            "error": format_validation_error(e),
            "today": today,
            "form_data": {
                "model": model, "firmware_version": firmware_version,
                "serial_number": serial_number, "box_number": box_number,
                "arrival_date": arrival_date
            }
        })

    # Проверка на дубликат
    existing = db.query(Terminal).filter(Terminal.serial_number == data.serial_number).first()
    if existing:
        return templates.TemplateResponse("terminal_add.html", {
            "request": request,
            "error": f"Терминал с серийным номером {data.serial_number} уже существует",
            "today": today,
            "form_data": {
                "model": model, "firmware_version": firmware_version,
                "serial_number": serial_number, "box_number": box_number,
                "arrival_date": arrival_date
            }
        })

    terminal = Terminal(
        model=data.model,
        firmware_version=data.firmware_version,
        serial_number=data.serial_number,
        box_number=data.box_number,
        arrival_date=data.arrival_date,
        status="warehouse"
    )
    db.add(terminal)
    db.flush()

    history = StatusHistory(
        terminal_id=terminal.id,
        old_status=None,
        new_status="warehouse",
        changed_by=user.id,
        comment="Терминал добавлен на склад"
    )
    db.add(history)
    db.commit()

    return RedirectResponse(url="/terminals", status_code=303)


# ============ МАССОВАЯ СМЕНА СТАТУСА ============
@router.post("/bulk-status")
async def bulk_change_status(
    request: Request,
    terminal_ids: str = Form(...),
    new_status: str = Form(...),
    bank: Optional[str] = Form(None),
    defect_type: Optional[str] = Form(None),
    defect_comment: Optional[str] = Form(None),
    comment: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user:
        return RedirectResponse(url="/auth/login", status_code=303)

    # Валидация статуса
    try:
        StatusChange(
            new_status=new_status,
            bank=bank.strip() if bank else None,
            defect_type=defect_type,
            defect_comment=defect_comment.strip() if defect_comment else None,
            comment=comment
        )
    except Exception as e:
        return HTMLResponse(f"Ошибка валидации: {format_validation_error(e)}", status_code=400)

    updated = 0
    for tid in terminal_ids.split(','):
        tid = tid.strip()
        if not tid:
            continue
        terminal = db.query(Terminal).filter(Terminal.id == int(tid)).first()
        if terminal:
            old_status = terminal.status
            old_bank = terminal.bank
            old_defect_type = terminal.defect_type

            terminal.status = new_status

            if new_status in ("reserved", "shipped"):
                terminal.bank = bank.strip() if bank else None
            else:
                terminal.bank = None

            if new_status in ("defective", "repair"):
                terminal.defect_type = defect_type
                terminal.defect_comment = defect_comment.strip() if defect_comment else None
            else:
                terminal.defect_type = None
                terminal.defect_comment = None

            history = StatusHistory(
                terminal_id=terminal.id,
                old_status=old_status,
                new_status=new_status,
                old_bank=old_bank,
                new_bank=terminal.bank,
                old_defect_type=old_defect_type,
                new_defect_type=terminal.defect_type,
                defect_comment=defect_comment if new_status in ("defective", "repair") else None,
                changed_by=user.id,
                comment=comment or "Массовая смена статуса"
            )
            db.add(history)
            updated += 1

    db.commit()
    return RedirectResponse(url="/terminals", status_code=303)


# ============ КАРТОЧКА ТЕРМИНАЛА ============
@router.get("/{terminal_id}", response_class=HTMLResponse)
async def terminal_card(
    terminal_id: int,
    request: Request,
    db: Session = Depends(get_db)
):
    terminal = db.query(Terminal).filter(Terminal.id == terminal_id).first()
    if not terminal:
        return HTMLResponse("Терминал не найден", status_code=404)

    history = db.query(StatusHistory).filter(
        StatusHistory.terminal_id == terminal_id
    ).order_by(StatusHistory.changed_at.desc()).all()

    return templates.TemplateResponse(
        "terminal_card.html",
        {"request": request, "terminal": terminal, "history": history}
    )


# ============ СМЕНА СТАТУСА В КАРТОЧКЕ ============
@router.post("/{terminal_id}/status")
async def terminal_change_status(
    terminal_id: int,
    request: Request,
    new_status: str = Form(...),
    bank: Optional[str] = Form(None),
    defect_type: Optional[str] = Form(None),
    defect_comment: Optional[str] = Form(None),
    comment: Optional[str] = Form(None),
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user:
        return RedirectResponse(url="/auth/login", status_code=303)

    terminal = db.query(Terminal).filter(Terminal.id == terminal_id).first()
    if not terminal:
        return HTMLResponse("Терминал не найден", status_code=404)

    # Валидация
    try:
        data = StatusChange(
            new_status=new_status,
            bank=bank.strip() if bank else None,
            defect_type=defect_type,
            defect_comment=defect_comment.strip() if defect_comment else None,
            comment=comment
        )
    except Exception as e:
        return HTMLResponse(f"Ошибка: {format_validation_error(e)}", status_code=400)

    old_status = terminal.status
    old_bank = terminal.bank
    old_defect_type = terminal.defect_type

    terminal.status = data.new_status

    if data.new_status in ("reserved", "shipped"):
        terminal.bank = data.bank
    else:
        terminal.bank = None

    if data.new_status in ("defective", "repair"):
        terminal.defect_type = data.defect_type
        terminal.defect_comment = data.defect_comment
    else:
        terminal.defect_type = None
        terminal.defect_comment = None

    history = StatusHistory(
        terminal_id=terminal.id,
        old_status=old_status,
        new_status=data.new_status,
        old_bank=old_bank,
        new_bank=terminal.bank,
        old_defect_type=old_defect_type,
        new_defect_type=terminal.defect_type,
        defect_comment=data.defect_comment if data.new_status in ("defective", "repair") else None,
        changed_by=user.id,
        comment=data.comment
    )
    db.add(history)
    db.commit()

    return RedirectResponse(url=f"/terminals/{terminal_id}", status_code=303)


# ============ РЕДАКТИРОВАНИЕ (только админ) ============
@router.get("/{terminal_id}/edit", response_class=HTMLResponse)
async def terminal_edit_form(
    terminal_id: int,
    request: Request,
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user or user.role != "admin":
        return HTMLResponse("Доступ запрещён. Требуются права администратора.", status_code=403)

    terminal = db.query(Terminal).filter(Terminal.id == terminal_id).first()
    if not terminal:
        return HTMLResponse("Терминал не найден", status_code=404)

    return templates.TemplateResponse(
        "terminal_edit.html",
        {"request": request, "terminal": terminal, "today": date.today().strftime('%Y-%m-%d')}
    )


@router.post("/{terminal_id}/edit")
async def terminal_edit(
    terminal_id: int,
    request: Request,
    model: str = Form(...),
    firmware_version: str = Form(...),
    serial_number: str = Form(...),
    box_number: str = Form(...),
    status: str = Form(...),
    bank: Optional[str] = Form(None),
    defect_type: Optional[str] = Form(None),
    defect_comment: Optional[str] = Form(None),
    arrival_date: str = Form(...),
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user or user.role != "admin":
        return HTMLResponse("Доступ запрещён. Требуются права администратора.", status_code=403)

    terminal = db.query(Terminal).filter(Terminal.id == terminal_id).first()
    if not terminal:
        return HTMLResponse("Терминал не найден", status_code=404)

    today = date.today().strftime('%Y-%m-%d')

    # Валидация
    try:
        data = TerminalUpdate(
            model=model.strip(),
            firmware_version=firmware_version.strip(),
            serial_number=serial_number.strip(),
            box_number=box_number.strip(),
            status=status,
            bank=bank.strip() if bank else None,
            defect_type=defect_type,
            defect_comment=defect_comment.strip() if defect_comment else None,
            arrival_date=arrival_date
        )
    except Exception as e:
        return templates.TemplateResponse("terminal_edit.html", {
            "request": request,
            "terminal": terminal,
            "today": today,
            "error": format_validation_error(e)
        })

    # Проверка на дубликат серийника
    existing = db.query(Terminal).filter(
        Terminal.serial_number == data.serial_number,
        Terminal.id != terminal_id
    ).first()
    if existing:
        return templates.TemplateResponse("terminal_edit.html", {
            "request": request,
            "terminal": terminal,
            "today": today,
            "error": f"Терминал с серийным номером {data.serial_number} уже существует"
        })

    old_status = terminal.status
    old_bank = terminal.bank
    old_defect_type = terminal.defect_type

    terminal.model = data.model
    terminal.firmware_version = data.firmware_version
    terminal.serial_number = data.serial_number
    terminal.box_number = data.box_number
    terminal.status = data.status
    terminal.bank = data.bank if data.status in ("reserved", "shipped") else None
    terminal.defect_type = data.defect_type if data.status in ("defective", "repair") else None
    terminal.defect_comment = data.defect_comment if data.status in ("defective", "repair") else None
    terminal.arrival_date = data.arrival_date

    if old_status != data.status or old_bank != terminal.bank or old_defect_type != terminal.defect_type:
        history = StatusHistory(
            terminal_id=terminal.id,
            old_status=old_status,
            new_status=data.status,
            old_bank=old_bank,
            new_bank=terminal.bank,
            old_defect_type=old_defect_type,
            new_defect_type=terminal.defect_type,
            defect_comment=data.defect_comment if data.status in ("defective", "repair") else None,
            changed_by=user.id,
            comment="Изменено администратором через редактирование"
        )
        db.add(history)

    db.commit()
    return RedirectResponse(url=f"/terminals/{terminal_id}", status_code=303)


# ============ УДАЛЕНИЕ (только админ) ============
@router.post("/{terminal_id}/delete")
async def terminal_delete(
    terminal_id: int,
    request: Request,
    db: Session = Depends(get_db)
):
    user = get_user_from_cookie(request, db)
    if not user or user.role != "admin":
        return HTMLResponse("Доступ запрещён. Требуются права администратора.", status_code=403)

    terminal = db.query(Terminal).filter(Terminal.id == terminal_id).first()
    if not terminal:
        return HTMLResponse("Терминал не найден", status_code=404)

    db.query(StatusHistory).filter(StatusHistory.terminal_id == terminal_id).delete()
    db.delete(terminal)
    db.commit()

    return RedirectResponse(url="/terminals", status_code=303)